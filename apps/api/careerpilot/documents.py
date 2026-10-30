"""Extractive document preparation. Every factual line retains current source references."""

import hashlib
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document as WordDocument
from docx.shared import Inches, Pt
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4, LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph, SimpleDocTemplate
from sqlalchemy import select

from .config import OUTPUTS
from .matching import BANNED, contains, fact_text, prefs, usable
from .models import FACTS, Candidate, Employer, Vacancy

TEMPLATES = {
    c: {
        "page": "Letter" if c in {"Canada", "USA"} else "A4",
        "date": "%m/%Y" if c == "Italy" else "%b %Y",
        "spelling": "US" if c == "USA" else "British",
        "nationality": c == "UAE",
        "columns": 1,
        "pages": 1,
        "header": "name | location | phone | email | LinkedIn | GitHub",
    }
    for c in [
        "Canada",
        "UAE",
        "USA",
        "United Kingdom",
        "Ireland",
        "Australia",
        "New Zealand",
        "Qatar",
        "Italy",
    ]
}
HEADINGS = {
    "SUMMARY",
    "EDUCATION",
    "EXPERIENCE",
    "PROJECTS / PUBLICATION",
    "SKILLS",
    "CERTIFICATIONS",
    "LANGUAGES",
}
FIXED = {"Dear Hiring Manager,", "Thank you for considering my application.", "Yours sincerely,"}
SPELLINGS = {
    "modelling": "modeling",
    "modeller": "modeler",
    "visualisation": "visualization",
    "organisation": "organization",
    "prioritise": "prioritize",
    "prioritised": "prioritized",
    "analyse": "analyze",
    "analysed": "analyzed",
    "optimisation": "optimization",
}


def spelling(value, country):
    pairs = SPELLINGS if country == "USA" else {v: k for k, v in SPELLINGS.items()}
    return re.sub(
        r"\b(" + "|".join(pairs) + r")\b",
        lambda m: pairs[m[0].lower()].capitalize() if m[0][0].isupper() else pairs[m[0].lower()],
        value,
        flags=re.I,
    )


def month(value, country):
    return datetime.fromisoformat(value + "-01" if len(value) == 7 else value).strftime(
        TEMPLATES[country]["date"]
    )


def sources(db):
    p = db.scalar(select(Candidate))
    result = {}
    if p.status == "confirmed" and p.allowed_on_resume:
        result[p.id] = {"name": p.name, "summary": p.summary, "headline": p.headline}
    for kind, model in FACTS.items():
        if kind == "rules":
            continue
        for f in db.scalars(select(model)):
            if usable(f):
                result[f.id] = {
                    k: getattr(f, k)
                    for k in [
                        "title",
                        "organization",
                        "details",
                        "start_date",
                        "end_date",
                        "learning",
                    ]
                }
    return result


def check(db, content, claims, country):
    current = sources(db)
    errors = []
    mapping = {c["line"]: c for c in claims}
    for index, line in enumerate(content.splitlines()):
        if BANNED.search(line):
            errors.append({"line": index + 1, "reason": "Prohibited content."})
        if contains(line, "AutoCAD") and not contains(line, "2D"):
            errors.append(
                {"line": index + 1, "reason": "AutoCAD must explicitly be limited to 2D."}
            )
        if line in HEADINGS:
            continue
        claim = mapping.get(index)
        if not claim:
            errors.append({"line": index + 1, "reason": "No provenance for this line."})
            continue
        if claim.get("type") == "fixed":
            if line not in FIXED:
                errors.append({"line": index + 1, "reason": "Unknown fixed text."})
            continue
        if claim.get("type") == "job":
            j = db.get(Vacancy, claim["id"])
            e = db.get(Employer, j.employer_id) if j else None
            if not j or line != f"I am applying for the {j.title} role at {e.name}.":
                errors.append({"line": index + 1, "reason": "Saved job changed."})
            continue
        if claim.get("type") == "contact":
            e = db.get(Employer, claim["id"])
            if not e or not any(
                c.get("verified") and c.get("source_url") and line == f"Dear {c.get('salutation')},"
                for c in e.contacts
            ):
                errors.append({"line": index + 1, "reason": "Contact is no longer confirmed."})
            continue
        values = []
        for piece in claim.get("pieces", []):
            source = current.get(piece["id"])
            field = piece["field"]
            value = piece["value"]
            if not source or field not in source:
                errors.append(
                    {"line": index + 1, "reason": "Source fact is excluded, removed or unverified."}
                )
                continue
            original = source[field]
            if field in {"start_date", "end_date"}:
                if not original or month(original, country) != value:
                    errors.append({"line": index + 1, "reason": "Date cannot be traced."})
            else:
                if piece.get("spelling"):
                    original = spelling(original, country)
                if value not in original and not (
                    piece.get("lower") and value and value[0].upper() + value[1:] in original
                ):
                    errors.append(
                        {
                            "line": index + 1,
                            "reason": "Claim cannot be traced to the confirmed source.",
                        }
                    )
            values.append(value)
        prefix = claim.get("prefix", "")
        suffix = claim.get("suffix", "")
        separator = claim.get("separator", "")
        if (
            prefix
            not in {
                "",
                "- ",
                "At ",
                "I ",
                "My work included ",
                "Nationality: ",
                "Expected ",
                "Ended ",
                "Started ",
            }
            or suffix not in {"", ":", " (currently learning)"}
            or separator not in {"", " | ", " to ", ", "}
        ):
            errors.append({"line": index + 1, "reason": "Unsupported claim formatting."})
        if not values or line != prefix + separator.join(values) + suffix:
            errors.append({"line": index + 1, "reason": "Untraceable text."})
    return {
        "passed": not errors,
        "errors": errors,
        "claims_checked": len(claims),
        "method": "Exact confirmed facts, explicit formatting and spelling transformations only.",
    }


def compose(db, job, kind, country, contact_index=None):
    if country not in TEMPLATES:
        raise ValueError("Choose a supported country template")
    p = db.scalar(select(Candidate))
    src = sources(db)
    pref = prefs(db)
    if p.id not in src or BANNED.search(p.name):
        raise ValueError("Confirm your name and allow it on resumes first")
    lines = []
    claims = []
    identity = {f.title.lower(): f for f in db.scalars(select(FACTS["identity"])) if f.id in src}
    scope = next(
        (f for f in db.scalars(select(FACTS["skills"])) if f.id in src and " (2D only)" in f.title),
        None,
    )

    def piece(f, field, value=None):
        return {"id": f.id, "field": field, "value": getattr(f, field) if value is None else value}

    def add(pieces, separator="", prefix="", suffix=""):
        pieces = [dict(v) for v in pieces if v["value"]]
        for v in pieces:
            if v["field"] in {"details", "summary"} and v["id"] not in {
                x.id for x in identity.values()
            }:
                changed = spelling(v["value"], country)
                if changed != v["value"]:
                    v.update(value=changed, spelling=True)
        text = prefix + separator.join(v["value"] for v in pieces) + suffix
        if contains(text, "AutoCAD") and not contains(text, "2D"):
            if not scope:
                return
            pieces.append(piece(scope, "title", " (2D only)"))
            text = prefix + separator.join(v["value"] for v in pieces) + suffix
        if not pieces or BANNED.search(text):
            return
        claims.append(
            {
                "line": len(lines),
                "pieces": pieces,
                "separator": separator,
                "prefix": prefix,
                "suffix": suffix,
            }
        )
        lines.append(text)

    def fixed(text):
        claims.append({"line": len(lines), "type": "fixed"})
        lines.append(text)

    add([piece(p, "name")])
    keys = ["location", "phone", "email", "linkedin", "github"]
    if country == "Canada" and pref["canadian_location"] and "canada location" in identity:
        keys[0] = "canada location"
    add([piece(identity[k], "details") for k in keys if k in identity], separator=" | ")
    if country == "UAE" and "nationality" in identity:
        add([piece(identity["nationality"], "details")], prefix="Nationality: ")
    if kind == "cover_letter":
        employer = db.get(Employer, job["employer_id"])
        contacts = employer.contacts
        contact = (
            contacts[contact_index]
            if contact_index is not None and contact_index < len(contacts)
            else None
        )
        if contact and contact.get("verified") and contact.get("salutation"):
            claims.append({"line": len(lines), "type": "contact", "id": employer.id})
            lines.append(f"Dear {contact['salutation']},")
        else:
            fixed("Dear Hiring Manager,")
        opening = f"I am applying for the {job['title']} role at {job['employer']}."
        if not BANNED.search(opening):
            claims.append({"line": len(lines), "type": "job", "id": job["id"]})
            lines.append(opening)
    if p.summary and not BANNED.search(p.summary):
        summary = p.summary.splitlines()[0]
        if len(summary) <= 260:
            if kind == "resume":
                lines.append("SUMMARY")
            add([piece(p, "summary", summary)])
    query = job["title"] + " " + job.get("description", "")

    def rank(value):
        return sum(
            contains(query, w)
            for w in set(re.findall(r"[\w-]{3,}", value.lower()))
            if w not in {"the", "and", "with", "for", "from"}
        )

    if kind == "cover_letter":
        for f in sorted(
            [f for f in db.scalars(select(FACTS["experience"])) if f.id in src],
            key=lambda f: (f.end_date, f.start_date),
            reverse=True,
        )[:2]:
            bullets = sorted(
                [
                    v.strip()
                    for v in f.details.splitlines()
                    if 0 < len(v.strip()) <= 320 and not BANNED.search(v)
                ],
                key=rank,
                reverse=True,
            )
            if not bullets:
                continue
            add([piece(f, "organization")], prefix="At ", suffix=":")
            value = bullets[0]
            v = piece(f, "details", value)
            first = value.split()[0].strip(".;,:")
            if not first.isupper():
                v.update(value=value[0].lower() + value[1:], lower=True)
            add(
                [v],
                prefix="I "
                if first.lower()
                in {
                    "built",
                    "traced",
                    "presented",
                    "prepared",
                    "supported",
                    "conducted",
                    "reviewed",
                    "performed",
                    "mapped",
                    "derived",
                    "rotated",
                }
                else "My work included ",
            )
        fixed("Thank you for considering my application.")
        fixed("Yours sincerely,")
        add([piece(p, "name")])
    else:
        groups = [
            ("education", "EDUCATION", 2),
            ("experience", "EXPERIENCE", 2),
            ("projects", "PROJECTS / PUBLICATION", 1),
            ("skills", "SKILLS", 8),
            ("certifications", "CERTIFICATIONS", 2),
        ] + ([("languages", "LANGUAGES", 3)] if pref["languages_on_resume"] else [])
        for category, title, limit in groups:
            records = sorted(
                [f for f in db.scalars(select(FACTS[category])) if f.id in src],
                key=(lambda f: (f.end_date, f.start_date))
                if category in {"education", "experience"}
                else lambda f: rank(fact_text(f)),
                reverse=True,
            )[:limit]
            if not records:
                continue
            lines.append(title)
            for f in records:
                add(
                    [piece(f, "title"), piece(f, "organization")],
                    separator=" | ",
                    suffix=" (currently learning)"
                    if f.learning and "currently learning" not in f.title.lower()
                    else "",
                )
                dates = [
                    piece(f, k, month(getattr(f, k), country))
                    for k in ["start_date", "end_date"]
                    if getattr(f, k)
                ]
                if dates and category in {"experience", "education"}:
                    prefix = (
                        "Expected "
                        if category == "education" and "expected" in f.details.lower()
                        else "Ended "
                        if category == "experience" and not f.start_date
                        else ""
                    )
                    add(dates, separator=" to ", prefix=prefix)
                bullets = sorted(
                    [
                        v.strip()
                        for v in f.details.splitlines()
                        if v.strip() and not BANNED.search(v)
                    ],
                    key=rank,
                    reverse=True,
                )
                if category in {"experience", "projects"}:
                    for v in [b for b in bullets if len(b) <= 320][
                        : 2 if category == "experience" else 1
                    ]:
                        add([piece(f, "details", v)], prefix="- ")
                elif category == "education":
                    for v in [b for b in bullets if len(b) <= 100 and "expected" not in b.lower()][
                        :1
                    ]:
                        add([piece(f, "details", v)])
    content = "\n".join(lines)
    checks = check(db, content, claims, country)
    if not checks["passed"]:
        raise ValueError("Claim checks failed. Review source facts before generating.")
    return content, claims, checks


def write_files(content, country, identifier):
    folder = OUTPUTS / identifier
    folder.mkdir(mode=0o700, parents=True, exist_ok=True)
    word = folder / "draft.docx"
    pdf = folder / "draft.pdf"
    lines = content.splitlines()
    doc = WordDocument()
    section = doc.sections[0]
    letter = TEMPLATES[country]["page"] == "Letter"
    section.page_width = Inches(8.5 if letter else 8.2677)
    section.page_height = Inches(11 if letter else 11.6929)
    section.top_margin = section.bottom_margin = Inches(0.52)
    section.left_margin = section.right_margin = Inches(0.6)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1
    for i, text in enumerate(lines):
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(text)
        if i == 0:
            run.bold = True
            run.font.size = Pt(15)
        elif text in HEADINGS:
            run.bold = True
            paragraph.paragraph_format.space_before = Pt(5)
    doc.save(word)
    conversion = "ReportLab fallback"
    converted = False
    if sys.platform in {"win32", "darwin"}:
        try:
            result = subprocess.run(
                [sys.executable, "-m", "careerpilot.word_convert", str(word), str(pdf)],
                capture_output=True,
                timeout=40,
            )
            converted = result.returncode == 0 and pdf.exists()
            if converted:
                conversion = "Microsoft Word via docx2pdf"
        except (OSError, subprocess.TimeoutExpired):
            pass
    office = shutil.which("libreoffice") or shutil.which("soffice")
    if not converted and office:
        try:
            with tempfile.TemporaryDirectory(prefix="careerpilot-office-") as temp:
                result = subprocess.run(
                    [
                        office,
                        "-env:UserInstallation=" + Path(temp).as_uri(),
                        "--headless",
                        "--convert-to",
                        "pdf",
                        "--outdir",
                        str(folder),
                        str(word),
                    ],
                    capture_output=True,
                    timeout=40,
                )
            converted = result.returncode == 0 and pdf.exists()
            if converted:
                conversion = "LibreOffice headless"
        except (OSError, subprocess.TimeoutExpired):
            pass
    if not converted:
        normal = ParagraphStyle("body", fontName="Helvetica", fontSize=9, leading=11, spaceAfter=2)
        bold = ParagraphStyle("head", parent=normal, fontName="Helvetica-Bold", spaceBefore=5)
        title = ParagraphStyle("title", parent=bold, fontSize=15, leading=18)
        SimpleDocTemplate(
            str(pdf),
            pagesize=LETTER if letter else A4,
            leftMargin=43,
            rightMargin=43,
            topMargin=38,
            bottomMargin=38,
        ).build(
            [
                Paragraph(escape(text), title if i == 0 else bold if text in HEADINGS else normal)
                for i, text in enumerate(lines)
            ]
        )
    if len(PdfReader(pdf).pages) != 1:
        raise ValueError(
            "Draft exceeds one page. Shorten confirmed bullets or disable less relevant facts, then generate again."
        )
    for f in [word, pdf]:
        f.chmod(0o600)
    return conversion, {
        "pages": 1,
        "docx_sha256": hashlib.sha256(word.read_bytes()).hexdigest(),
        "pdf_sha256": hashlib.sha256(pdf.read_bytes()).hexdigest(),
    }
